import re
import os
import io
import csv
import datetime
from flask import Blueprint, request, jsonify, send_file, session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.db import conectar_db, resolver_listas_precios, detectar_columna_equivalencia, obtener_siguiente_id, conectar_sqlite
from app.services.compras_service import normalizar_clave_ceros

compras_bp = Blueprint('compras_bp', __name__)

# ================= MÓDULO 1: VERIFICAR Y ACTUALIZAR ARTÍCULOS =================

@compras_bp.route('/api/verificar-articulos', methods=['POST'])
def verificar_articulos():
    data = request.get_json() or {}
    claves = data.get('claves', [])
    if not claves:
        return jsonify({"existentes": {}, "claves_existentes": []})

    conn = None
    try:
        conn = conectar_db()
        cursor = conn.cursor()
        listas_map = resolver_listas_precios(conn)
        id_a_lista = {v: k for k, v in listas_map.items()}
        col_equiv = detectar_columna_equivalencia(cursor)

        claves_limpias = list({str(c).strip().upper() for c in claves if str(c).strip()})
        mapa_variantes = {}  # variante_a_buscar -> set de claves solicitadas
        for c in claves_limpias:
            mapa_variantes.setdefault(c, set()).add(c)
            c_norm = normalizar_clave_ceros(c)
            if c_norm != c:
                mapa_variantes.setdefault(c_norm, set()).add(c)

        todas_claves = list(mapa_variantes.keys())
        mapa_existentes = {}

        lote_size = 100
        for i in range(0, len(todas_claves), lote_size):
            lote = todas_claves[i:i + lote_size]
            placeholders = ','.join(['?'] * len(lote))
            
            select_equiv = f"TRIM(la.{col_equiv})" if col_equiv else "''"
            join_equiv = "LEFT JOIN LIBRES_ARTICULOS la ON la.ARTICULO_ID = a.ARTICULO_ID" if col_equiv else ""

            query = f"""
                SELECT 
                    TRIM(ca.CLAVE_ARTICULO),
                    COALESCE(NULLIF(TRIM(csec.CLAVE), ''), l.NOMBRE) AS CLAVE_LINEA,
                    a.ARTICULO_ID,
                    {select_equiv} AS EQUIVALENCIA_ACTUAL,
                    TRIM(a.NOMBRE) AS NOMBRE_OFICIAL_MICROSIP
                FROM CLAVES_ARTICULOS ca
                JOIN ARTICULOS a ON a.ARTICULO_ID = ca.ARTICULO_ID
                LEFT JOIN LINEAS_ARTICULOS l ON l.LINEA_ARTICULO_ID = a.LINEA_ARTICULO_ID
                LEFT JOIN CLAVES_CAT_SEC csec ON csec.ELEM_ID = l.LINEA_ARTICULO_ID AND csec.NOMBRE_TABLA = 'LINEAS_ARTICULOS'
                {join_equiv}
                WHERE TRIM(ca.CLAVE_ARTICULO) IN ({placeholders})
            """
            cursor.execute(query, lote)
            filas_art = cursor.fetchall()
            
            art_ids_map = {}  # art_id -> list of target keys to receive prices
            for row in filas_art:
                if row[0]:
                    clave_art = str(row[0]).strip().upper()
                    clave_lin = str(row[1]).strip() if row[1] else ""
                    art_id = int(row[2])
                    equiv_act = str(row[3]).strip() if row[3] else ""
                    nombre_microsip = str(row[4]).strip() if len(row) > 4 and row[4] else ""

                    claves_dest = set(mapa_variantes.get(clave_art, [clave_art]))
                    claves_dest.add(clave_art)

                    if art_id not in art_ids_map:
                        art_ids_map[art_id] = []

                    for k in claves_dest:
                        if k not in art_ids_map[art_id]:
                            art_ids_map[art_id].append(k)
                    mapa_existentes[k] = {
                        "clave_microsip": clave_art,
                        "linea": clave_lin,
                        "equivalencia": equiv_act,
                        "nombre_microsip": nombre_microsip,
                        "costo_bd": None,
                        "costo_bd_con_iva": None,
                        "precios": {"publico": None, "talleres": None, "mayoreo": None, "minimo": None},
                        "precios_sin_iva": {"publico": None, "talleres": None, "mayoreo": None, "minimo": None},
                        "tasa_iva": 16.0
                    }

            if art_ids_map:
                art_id_placeholders = ','.join(['?'] * len(art_ids_map))
                art_id_keys = list(art_ids_map.keys())

                # Consultar tasa de IVA de cada artículo
                impuestos_art = {}
                try:
                    q_tax = f"""
                        SELECT ia.ARTICULO_ID, COALESCE(i.PCTJE_IMPUESTO, 16.0)
                        FROM IMPUESTOS_ARTICULOS ia
                        JOIN IMPUESTOS i ON i.IMPUESTO_ID = ia.IMPUESTO_ID
                        WHERE ia.ARTICULO_ID IN ({art_id_placeholders})
                    """
                    cursor.execute(q_tax, art_id_keys)
                    for r_tax in cursor.fetchall():
                        aid = int(r_tax[0])
                        pct = float(r_tax[1]) if r_tax[1] is not None else 16.0
                        impuestos_art[aid] = pct
                except Exception as e_tax:
                    print("Aviso al consultar IMPUESTOS_ARTICULOS:", e_tax)

                q_precios = f"""
                    SELECT ARTICULO_ID, PRECIO_EMPRESA_ID, PRECIO
                    FROM PRECIOS_ARTICULOS
                    WHERE ARTICULO_ID IN ({art_id_placeholders})
                """
                cursor.execute(q_precios, art_id_keys)
                for p_art_id, p_emp_id, precio_val in cursor.fetchall():
                    if p_art_id in art_ids_map and p_emp_id in id_a_lista:
                        tipo_lista = id_a_lista[p_emp_id]
                        val_num = float(precio_val) if precio_val is not None else None
                        tax_pct = impuestos_art.get(p_art_id, 16.0)
                        factor_iva = 1.0 + (tax_pct / 100.0)
                        val_con_iva = round(val_num * factor_iva, 4) if val_num is not None else None

                        for k in art_ids_map[p_art_id]:
                            if k in mapa_existentes:
                                # Precios con impuesto para comparativa exacta vs sugerido (c/IVA)
                                mapa_existentes[k]["precios"][tipo_lista] = val_con_iva
                                mapa_existentes[k]["precios_sin_iva"][tipo_lista] = val_num
                                mapa_existentes[k]["tasa_iva"] = tax_pct

                # Consultar último costo de compra en BD para comparar Factura vs BD con IVA
                costos_bd = {}
                try:
                    q_cm = f"""
                        SELECT d.ARTICULO_ID, d.PRECIO_UNITARIO
                        FROM DOCTOS_CM_DET d
                        JOIN DOCTOS_CM cm ON cm.DOCTO_CM_ID = d.DOCTO_CM_ID
                        WHERE d.ARTICULO_ID IN ({art_id_placeholders}) AND cm.TIPO_DOCTO IN ('C', 'R')
                        ORDER BY d.DOCTO_CM_ID DESC
                    """
                    cursor.execute(q_cm, art_id_keys)
                    for r_cm in cursor.fetchall():
                        aid = int(r_cm[0])
                        if aid not in costos_bd and r_cm[1] is not None and float(r_cm[1]) > 0:
                            costos_bd[aid] = float(r_cm[1])
                except Exception as e_cost:
                    print("Aviso al consultar DOCTOS_CM_DET en compras:", e_cost)

                ids_sin_costo = [aid for aid in art_id_keys if aid not in costos_bd]
                if ids_sin_costo:
                    try:
                        p_sin = ','.join(['?'] * len(ids_sin_costo))
                        q_pc = f"""
                            SELECT pc.ARTICULO_ID, pcd.PRECIO_UVEN
                            FROM PRECIOS_COMPRA pc
                            JOIN PRECIOS_COMPRA_DET pcd ON pcd.PRECIO_COMPRA_ID = pc.PRECIO_COMPRA_ID
                            WHERE pc.ARTICULO_ID IN ({p_sin}) AND pcd.PRECIO_UVEN > 0
                        """
                        cursor.execute(q_pc, ids_sin_costo)
                        for r_pc in cursor.fetchall():
                            aid = int(r_pc[0])
                            if aid not in costos_bd and r_pc[1] is not None and float(r_pc[1]) > 0:
                                costos_bd[aid] = float(r_pc[1])
                    except Exception as e_pc:
                        print("Aviso al consultar PRECIOS_COMPRA en compras:", e_pc)

                for aid, c_val in costos_bd.items():
                    if aid in art_ids_map:
                        tax_pct = impuestos_art.get(aid, 16.0)
                        factor_iva = 1.0 + (tax_pct / 100.0)
                        for k in art_ids_map[aid]:
                            if k in mapa_existentes:
                                mapa_existentes[k]["costo_bd"] = round(c_val, 2)
                                mapa_existentes[k]["costo_bd_con_iva"] = round(c_val * factor_iva, 2)

        cursor.close()
        conn.close()
        return jsonify({"existentes": mapa_existentes, "claves_existentes": list(mapa_existentes.keys())})
    except Exception as e:
        if conn:
            conn.close()
        return jsonify({"error": str(e)}), 500

@compras_bp.route('/api/buscar-articulos-equivalencia', methods=['GET'])
def buscar_articulos_equivalencia():
    termino = request.args.get('q', '').strip().upper()
    if not termino or len(termino) < 2:
        return jsonify([])

    conn = None
    try:
        conn = conectar_db()
        cursor = conn.cursor()
        param = f"%{termino}%"
        termino_norm = normalizar_clave_ceros(termino)
        if termino_norm and termino_norm != termino:
            param_norm = f"%{termino_norm}%"
            query = """
                SELECT FIRST 20
                    TRIM(ca.CLAVE_ARTICULO) AS CLAVE,
                    TRIM(a.NOMBRE) AS NOMBRE
                FROM ARTICULOS a
                JOIN CLAVES_ARTICULOS ca ON ca.ARTICULO_ID = a.ARTICULO_ID
                WHERE UPPER(ca.CLAVE_ARTICULO) LIKE ? OR UPPER(a.NOMBRE) LIKE ?
                   OR UPPER(ca.CLAVE_ARTICULO) LIKE ? OR UPPER(a.NOMBRE) LIKE ?
                ORDER BY a.NOMBRE
            """
            cursor.execute(query, (param, param, param_norm, param_norm))
        else:
            query = """
                SELECT FIRST 20
                    TRIM(ca.CLAVE_ARTICULO) AS CLAVE,
                    TRIM(a.NOMBRE) AS NOMBRE
                FROM ARTICULOS a
                JOIN CLAVES_ARTICULOS ca ON ca.ARTICULO_ID = a.ARTICULO_ID
                WHERE UPPER(ca.CLAVE_ARTICULO) LIKE ? OR UPPER(a.NOMBRE) LIKE ?
                ORDER BY a.NOMBRE
            """
            cursor.execute(query, (param, param))
        resultados = [{"clave": str(row[0]).strip() if row[0] else "", "nombre": str(row[1]).strip() if row[1] else ""} for row in cursor.fetchall()]
        cursor.close()
        conn.close()
        return jsonify(resultados)
    except Exception as e:
        if conn:
            conn.close()
        return jsonify({"error": str(e)}), 500

@compras_bp.route('/api/actualizar-equivalencias', methods=['POST'])
def actualizar_equivalencias():
    data = request.get_json() or {}
    articulos = data.get('articulos', [])
    if not articulos:
        return jsonify({"error": "No se recibieron equivalencias para procesar."}), 400

    conn = None
    actualizados = 0
    insertados = 0
    no_encontrados = []

    try:
        conn = conectar_db()
        cursor = conn.cursor()
        col_equiv = detectar_columna_equivalencia(cursor)

        if not col_equiv:
            cursor.close()
            conn.close()
            return jsonify({"error": "No se encontró el campo 'Equivalencia1' en la tabla LIBRES_ARTICULOS."}), 500

        for art in articulos:
            clave_art = str(art.get('clave', '')).strip().upper()
            codigo_equiv = str(art.get('equivalencia', '')).strip().upper()
            if not clave_art or not codigo_equiv:
                continue

            cursor.execute("SELECT FIRST 1 ARTICULO_ID FROM CLAVES_ARTICULOS WHERE TRIM(CLAVE_ARTICULO) = ?", (clave_art,))
            row = cursor.fetchone()
            if not row or not row[0]:
                c_norm = normalizar_clave_ceros(clave_art)
                if c_norm != clave_art:
                    cursor.execute("SELECT FIRST 1 ARTICULO_ID FROM CLAVES_ARTICULOS WHERE TRIM(CLAVE_ARTICULO) = ?", (c_norm,))
                    row = cursor.fetchone()

            if not row or not row[0]:
                no_encontrados.append(clave_art)
                continue

            articulo_id = int(row[0])
            cursor.execute("SELECT FIRST 1 ARTICULO_ID FROM LIBRES_ARTICULOS WHERE ARTICULO_ID = ?", (articulo_id,))
            if cursor.fetchone():
                cursor.execute(f"UPDATE LIBRES_ARTICULOS SET {col_equiv} = ? WHERE ARTICULO_ID = ?", (codigo_equiv, articulo_id))
                actualizados += 1
            else:
                cursor.execute(f"INSERT INTO LIBRES_ARTICULOS (ARTICULO_ID, {col_equiv}) VALUES (?, ?)", (articulo_id, codigo_equiv))
                insertados += 1

        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"success": True, "actualizados": actualizados, "insertados": insertados, "no_encontrados": no_encontrados, "campo_usado": col_equiv})
    except Exception as e:
        if conn:
            try:
                conn.rollback()
            except:
                pass
            conn.close()
        return jsonify({"error": str(e), "success": False}), 500

@compras_bp.route('/api/actualizar-precios', methods=['POST'])
def actualizar_precios():
    data = request.get_json() or {}
    articulos = data.get('articulos', [])
    if not articulos:
        return jsonify({"error": "No se enviaron artículos para procesar."}), 400

    conn = None
    actualizados = 0
    insertados = 0
    omitidos_no_subieron = 0
    no_encontrados = []
    
    try:
        conn = conectar_db()
        listas_map = resolver_listas_precios(conn)
        cursor = conn.cursor()
        
        for art in articulos:
            clave = str(art.get('clave', '')).strip().upper()
            if not clave:
                continue

            cursor.execute("SELECT FIRST 1 ARTICULO_ID FROM CLAVES_ARTICULOS WHERE TRIM(CLAVE_ARTICULO) = ?", (clave,))
            row = cursor.fetchone()
            if not row or not row[0]:
                c_norm = normalizar_clave_ceros(clave)
                if c_norm != clave:
                    cursor.execute("SELECT FIRST 1 ARTICULO_ID FROM CLAVES_ARTICULOS WHERE TRIM(CLAVE_ARTICULO) = ?", (c_norm,))
                    row = cursor.fetchone()

            if not row or not row[0]:
                no_encontrados.append(clave)
                continue
                
            articulo_id = int(row[0])

            # Obtener tasa de impuesto del artículo
            tax_pct = 16.0
            try:
                cursor.execute("""
                    SELECT FIRST 1 COALESCE(i.PCTJE_IMPUESTO, 16.0)
                    FROM IMPUESTOS_ARTICULOS ia
                    JOIN IMPUESTOS i ON i.IMPUESTO_ID = ia.IMPUESTO_ID
                    WHERE ia.ARTICULO_ID = ?
                """, (articulo_id,))
                r_tax = cursor.fetchone()
                if r_tax and r_tax[0] is not None:
                    tax_pct = float(r_tax[0])
            except:
                tax_pct = 16.0

            factor_iva = 1.0 + (tax_pct / 100.0)

            precios_map = {
                listas_map["publico"]: round(float(art.get('precioPublico', 0)), 2),
                listas_map["talleres"]: round(float(art.get('precioTalleres', 0)), 2),
                listas_map["mayoreo"]: round(float(art.get('precioMayoreo', 0)), 2),
                listas_map["minimo"]: round(float(art.get('precioMinimo', 0)), 2)
            }

            for lista_id, precio_val in precios_map.items():
                if precio_val <= 0:
                    continue

                cursor.execute(
                    "SELECT FIRST 1 PRECIO_ARTICULO_ID, PRECIO FROM PRECIOS_ARTICULOS WHERE ARTICULO_ID = ? AND PRECIO_EMPRESA_ID = ?",
                    (articulo_id, lista_id)
                )
                p_row = cursor.fetchone()

                # precio_val viene con IVA; PRECIOS_ARTICULOS.PRECIO en Microsip se almacena como precio neto de venta
                precio_neto_guardar = round(precio_val / factor_iva, 2)

                if p_row and p_row[0]:
                    precio_actual_neto = float(p_row[1]) if p_row[1] is not None else 0.0
                    precio_actual_con_iva = round(precio_actual_neto * factor_iva, 4)

                    # Comparar con IVA para saber si realmente aumentó el precio al cliente
                    if (precio_val - precio_actual_con_iva) > 0.05:
                        precio_art_id = int(p_row[0])
                        cursor.execute("""
                            UPDATE PRECIOS_ARTICULOS 
                            SET PRECIO = ?, FECHA_HORA_ULT_MODIF = CURRENT_TIMESTAMP
                            WHERE PRECIO_ARTICULO_ID = ?
                        """, (precio_neto_guardar, precio_art_id))
                        actualizados += 1
                    else:
                        omitidos_no_subieron += 1
                else:
                    nuevo_id = obtener_siguiente_id(cursor, "GEN_PRECIOS_ARTICULOS", "PRECIOS_ARTICULOS", "PRECIO_ARTICULO_ID")
                    cursor.execute("""
                        INSERT INTO PRECIOS_ARTICULOS 
                        (PRECIO_ARTICULO_ID, ARTICULO_ID, PRECIO_EMPRESA_ID, PRECIO, MONEDA_ID, MARGEN, MARKUP, FECHA_HORA_ULT_MODIF)
                        VALUES (?, ?, ?, ?, 1, 0, 0, CURRENT_TIMESTAMP)
                    """, (nuevo_id, articulo_id, lista_id, precio_neto_guardar))
                    insertados += 1

        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"success": True, "actualizados": actualizados, "insertados": insertados, "omitidos_no_subieron": omitidos_no_subieron, "no_encontrados": no_encontrados})
    except Exception as e:
        if conn:
            try:
                conn.rollback()
            except:
                pass
            conn.close()
        return jsonify({"error": str(e), "success": False}), 500

# ================= MÓDULO 2: EXCEL / CSV =================

@compras_bp.route('/api/exportar-excel-nuevos', methods=['POST'])
def exportar_excel_nuevos():
    if not openpyxl:
        return jsonify({"error": "openpyxl no instalado"}), 500

    data = request.get_json() or {}
    articulos = data.get('articulos', [])
    proveedor = str(data.get('proveedor', 'NIKKO')).strip() or 'NIKKO'

    if not articulos:
        return jsonify({"error": "No hay artículos nuevos para exportar"}), 400

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NUEVO"

    ws.cell(row=2, column=1, value=proveedor).font = Font(bold=True, size=14)
    headers = [
        "Clave", "Nombre Articulo", "Costo Unitario",
        "PRECIO PUBLICO CON IVA", "PRECIO TALLERES CON IVA", "PRECIO MAYOREO CON IVA", "PRECIO MINIMO CON IVA",
        None,
        "MARGEN DE DIFERENCIA\nPUBLICO", "MARGEN DE DIFERENCIA\nTALLERES", "MARGEN DE DIFERENCIA\nMAYOREO", "MARGEN DE DIFERENCIA\nMINIMO",
        None,
        "PUBLICO", "TALLERES", "MAYOREO", "MINIMO"
    ]

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill_blue = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_fill_green = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
    header_fill_dark = PatternFill(start_color="262626", end_color="262626", fill_type="solid")

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        if h:
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if col_idx in [4, 5, 6, 7]:
                cell.fill = header_fill_blue
            elif col_idx in [9, 10, 11, 12]:
                cell.fill = header_fill_green
            elif col_idx in [14, 15, 16, 17]:
                cell.fill = header_fill_dark

    row_num = 4
    for art in articulos:
        clave = str(art.get('codigo', '')).strip()
        nombre = str(art.get('descripcion', '')).strip()
        costo = float(art.get('costo', 0.0))

        ws.cell(row=row_num, column=1, value=clave)
        ws.cell(row=row_num, column=2, value=nombre)
        ws.cell(row=row_num, column=3, value=costo).number_format = '"$"#,##0.00'

        for c_idx, col_let in [(4, 'N'), (5, 'O'), (6, 'P'), (7, 'Q')]:
            ws.cell(row=row_num, column=c_idx, value=f"=(C{row_num}/((100-{col_let}{row_num})/100))*1.16").number_format = '"$"#,##0.00'

        for m_idx, col_let in [(9, 'D'), (10, 'E'), (11, 'F'), (12, 'G')]:
            ws.cell(row=row_num, column=m_idx, value=f"=((C{row_num}/({col_let}{row_num}/1.16))-1)*-1").number_format = '0.00%'

        ws.cell(row=row_num, column=14, value=60)
        ws.cell(row=row_num, column=15, value=50)
        ws.cell(row=row_num, column=16, value=40)
        ws.cell(row=row_num, column=17, value=30)
        row_num += 1

    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"PRECIOS_NUEVOS_{proveedor.replace(' ', '_')}.xlsx")

@compras_bp.route('/api/procesar-excel-precios', methods=['POST'])
def procesar_excel_precios():
    if 'file' not in request.files:
        return jsonify({"error": "No se subió ningún archivo"}), 400

    file = request.files['file']
    filename = file.filename.lower()
    items = []
    try:
        if filename.endswith('.csv'):
            stream = io.StringIO(file.stream.read().decode("utf-8", errors="ignore"))
            reader = csv.reader(stream)
            headers = [h.strip().upper() for h in next(reader, [])]
            idx_clave = next((i for i, h in enumerate(headers) if any(k in h for k in ['CLAVE', 'CODIGO', 'ARTICULO', 'PARTE'])), 0)
            idx_costo = next((i for i, h in enumerate(headers) if any(k in h for k in ['COSTO', 'PRECIO', 'VALOR', 'NUEVO'])), 1)
            for row in reader:
                if len(row) > max(idx_clave, idx_costo):
                    c = str(row[idx_clave]).strip().upper()
                    try:
                        costo = float(str(row[idx_costo]).replace('$', '').replace(',', '').strip())
                    except:
                        costo = 0.0
                    if c and costo > 0:
                        items.append({"clave": c, "costo": costo})

        elif filename.endswith(('.xlsx', '.xls')):
            if not openpyxl:
                return jsonify({"error": "openpyxl no instalado"}), 500
            wb = openpyxl.load_workbook(file, data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                return jsonify({"error": "El archivo está vacío"}), 400
            headers = [str(h).strip().upper() if h else '' for h in rows[0]]
            idx_clave = next((i for i, h in enumerate(headers) if any(k in h for k in ['CLAVE', 'CODIGO', 'ARTICULO', 'PARTE'])), 0)
            idx_costo = next((i for i, h in enumerate(headers) if any(k in h for k in ['COSTO', 'PRECIO', 'VALOR', 'NUEVO'])), 1)
            for r in rows[1:]:
                if r and len(r) > max(idx_clave, idx_costo):
                    c = str(r[idx_clave]).strip().upper() if r[idx_clave] is not None else ''
                    try:
                        costo = float(r[idx_costo]) if r[idx_costo] is not None else 0.0
                    except:
                        costo = 0.0
                    if c and costo > 0:
                        items.append({"clave": c, "costo": costo})
        else:
            return jsonify({"error": "Formato de archivo no soportado"}), 400
    except Exception as e:
        return jsonify({"error": f"Error al leer archivo: {str(e)}"}), 500

    return jsonify({"articulos": items, "total": len(items)})

# ================= GENERACIÓN DE EXCEL DE FACTURAS (E:\) =================

@compras_bp.route('/api/compras/verificar-excel-factura', methods=['GET'])
def verificar_excel_factura():
    try:
        folio = str(request.args.get('folio', '')).strip()
        if not folio:
            return jsonify({"existe": False})
        folio_clean = re.sub(r'[^\w\-]', '_', folio)
        nombre_archivo = f"{folio_clean}.xlsx"
        ruta_e = os.path.join("E:\\", nombre_archivo)
        existe = os.path.exists(ruta_e)
        fecha_mod = None
        tamano_kb = None
        if existe:
            stats = os.stat(ruta_e)
            fecha_mod = datetime.datetime.fromtimestamp(stats.st_mtime).strftime("%d/%m/%Y %I:%M %p")
            tamano_kb = round(stats.st_size / 1024, 1)
        return jsonify({
            "existe": existe,
            "nombre_archivo": nombre_archivo,
            "ruta_e": ruta_e,
            "fecha_mod": fecha_mod,
            "tamano_kb": tamano_kb
        })
    except Exception as e:
        return jsonify({"existe": False, "error": str(e)})

@compras_bp.route('/api/compras/generar-excel-factura', methods=['POST'])
def generar_excel_factura():
    try:
        data = request.get_json() or {}
        folio = str(data.get('folio', '')).strip()
        partidas = data.get('partidas', [])
        sobrescribir = bool(data.get('sobrescribir', False))

        if not folio:
            folio = "FACTURA_" + datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

        folio_clean = re.sub(r'[^\w\-]', '_', folio)
        nombre_archivo = f"{folio_clean}.xlsx"
        ruta_e = os.path.join("E:\\", nombre_archivo)

        # Si ya existe en E:\ y el usuario no indicó expresamente sobrescribir, avisar estilo Windows
        if os.path.exists(ruta_e) and not sobrescribir:
            stats = os.stat(ruta_e)
            fecha_mod = datetime.datetime.fromtimestamp(stats.st_mtime).strftime("%d/%m/%Y %I:%M %p")
            tamano_kb = round(stats.st_size / 1024, 1)
            return jsonify({
                "success": False,
                "ya_existe": True,
                "nombre_archivo": nombre_archivo,
                "ruta_e": ruta_e,
                "fecha_mod": fecha_mod,
                "tamano_kb": tamano_kb,
                "mensaje": f"{ruta_e} ya existe.\n¿Desea reemplazarlo?"
            })

        if not openpyxl:
            return jsonify({"success": False, "error": "openpyxl no está instalado en el servidor"}), 500

        wb = openpyxl.Workbook()
        ws2 = wb.active
        ws2.title = "Hoja2"
        ws1 = wb.create_sheet(title="Hoja1")
        wb.active = ws2

        # Encabezados exactos como en el formato de la empresa
        headers = ["Clave_Articulo", "Nombre_Articulo", "Unidades", "Precio_Unitario"]
        ws2.append(headers)

        font_header = Font(name="Calibri", size=11, bold=True)
        for col_idx in range(1, 5):
            cell = ws2.cell(row=1, column=col_idx)
            cell.font = font_header

        for p in partidas:
            clave_val = p.get('clave') or p.get('codigo') or ''
            if str(clave_val).isdigit():
                clave_val = int(clave_val)
            else:
                clave_val = str(clave_val).strip()

            nombre_val = str(p.get('nombre') or p.get('descripcion') or '').strip().upper()

            try:
                unidades_val = float(p.get('unidades') or p.get('cantidad') or 1)
                if unidades_val.is_integer():
                    unidades_val = int(unidades_val)
            except:
                unidades_val = 1

            try:
                precio_val = float(p.get('precio') or p.get('costo') or 0.0)
            except:
                precio_val = 0.0

            ws2.append([clave_val, nombre_val, unidades_val, precio_val])
            row_idx = ws2.max_row
            ws2.cell(row=row_idx, column=4).number_format = '#,##0.00'

        # Ajuste de ancho de columnas
        ws2.column_dimensions['A'].width = 18
        ws2.column_dimensions['B'].width = 40
        ws2.column_dimensions['C'].width = 12
        ws2.column_dimensions['D'].width = 18

        guardado_en_e = False
        error_guardado_e = None

        try:
            wb.save(ruta_e)
            guardado_en_e = True
        except Exception as e_save:
            error_guardado_e = str(e_save)
            print(f"Aviso al guardar en E:\\: {e_save}")

        return jsonify({
            "success": True,
            "sobrescrito": sobrescribir,
            "folio": folio_clean,
            "nombre_archivo": nombre_archivo,
            "ruta_e": ruta_e,
            "guardado_en_e": guardado_en_e,
            "error_e": error_guardado_e,
            "total_partidas": len(partidas),
            "mensaje": f"Archivo {'reemplazado / sobrescrito' if sobrescribir else 'guardado'} exitosamente en E:\\{nombre_archivo}" if guardado_en_e else f"No se pudo guardar en E:\\: {error_guardado_e}"
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@compras_bp.route('/api/compras/descargar-excel-factura/<path:nombre_archivo>', methods=['GET'])
def descargar_excel_factura(nombre_archivo):
    try:
        ruta = os.path.join("E:\\", nombre_archivo)
        if os.path.exists(ruta):
            return send_file(ruta, as_attachment=True, download_name=nombre_archivo)
        return jsonify({"error": "Archivo no encontrado"}), 404
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ================= MÓDULO 6: VALIDACIÓN Y AUTORIZACIÓN DE SOLICITUDES DE TRASPASO (PV) =================

@compras_bp.route('/api/compras/solicitudes-traspasos/conteo-pendientes', methods=['GET'])
def get_conteo_solicitudes_pendientes():
    try:
        conn = conectar_sqlite()
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM solicitudes_traspasos WHERE estatus = 'PENDIENTE_COMPRAS'")
        conteo = cur.fetchone()[0] or 0
        conn.close()
        return jsonify({"success": True, "pendientes": conteo})
    except Exception as e:
        return jsonify({"success": False, "pendientes": 0, "error": str(e)}), 500

@compras_bp.route('/api/compras/solicitudes-traspasos', methods=['GET'])
def get_solicitudes_traspasos():
    estatus_filtro = request.args.get('estatus', '').strip().upper()
    conn_sq = None
    try:
        conn_sq = conectar_sqlite()
        cur_sq = conn_sq.cursor()

        if estatus_filtro and estatus_filtro != 'TODAS':
            cur_sq.execute("""
                SELECT id, folio, fecha_solicitud, sucursal_destino_id, sucursal_destino_nombre,
                       articulo_id, clave, nombre, cantidad, stock_cedis_al_solicitar,
                       stock_local_al_solicitar, cliente_nombre, notas, usuario_solicita,
                       vendedor_nombre, estatus, usuario_autoriza, fecha_autorizacion,
                       motivo_rechazo, folio_traspaso_generado
                FROM solicitudes_traspasos
                WHERE estatus = ?
                ORDER BY id DESC LIMIT 200
            """, (estatus_filtro,))
        else:
            cur_sq.execute("""
                SELECT id, folio, fecha_solicitud, sucursal_destino_id, sucursal_destino_nombre,
                       articulo_id, clave, nombre, cantidad, stock_cedis_al_solicitar,
                       stock_local_al_solicitar, cliente_nombre, notas, usuario_solicita,
                       vendedor_nombre, estatus, usuario_autoriza, fecha_autorizacion,
                       motivo_rechazo, folio_traspaso_generado
                FROM solicitudes_traspasos
                ORDER BY CASE WHEN estatus = 'PENDIENTE_COMPRAS' THEN 0 ELSE 1 END, id DESC LIMIT 200
            """)

        filas = cur_sq.fetchall()
        conn_sq.close()

        # Consultar existencias EN VIVO en CEDIS para los artículos recuperados
        art_ids = list({int(r[5]) for r in filas if r[5]})
        stock_cedis_vivo = {}
        stock_dest_vivo = {}

        if art_ids:
            try:
                conn_fb = conectar_db()
                cur_fb = conn_fb.cursor()
                placeholders = ",".join("?" for _ in art_ids)
                cur_fb.execute(f"""
                    SELECT ARTICULO_ID, ALMACEN_ID, SUM(ENTRADAS_UNIDADES - SALIDAS_UNIDADES)
                    FROM SALDOS_IN
                    WHERE ARTICULO_ID IN ({placeholders})
                    GROUP BY ARTICULO_ID, ALMACEN_ID
                """, art_ids)
                for r_fb in cur_fb.fetchall():
                    a_id = int(r_fb[0])
                    alm_id = int(r_fb[1]) if r_fb[1] is not None else 0
                    cant = max(0, int(float(r_fb[2] or 0)))
                    if alm_id == 620110:
                        stock_cedis_vivo[a_id] = cant
                    else:
                        stock_dest_vivo.setdefault(a_id, {})[alm_id] = cant
                cur_fb.close()
                conn_fb.close()
            except Exception as e_fb:
                print("Aviso al consultar stock en vivo en Firebird para compras:", e_fb)

        solicitudes = []
        for r in filas:
            aid = int(r[5])
            dest_id = int(r[3])
            stk_c_vivo = stock_cedis_vivo.get(aid, float(r[9] or 0))
            stk_d_vivo = stock_dest_vivo.get(aid, {}).get(dest_id, float(r[10] or 0))

            solicitudes.append({
                "id": r[0],
                "folio": r[1],
                "fecha": r[2],
                "sucursal_destino_id": dest_id,
                "sucursal_destino": r[4],
                "articulo_id": aid,
                "clave": r[6],
                "nombre": r[7],
                "cantidad": float(r[8] or 0),
                "stock_cedis_inicial": float(r[9] or 0),
                "stock_cedis_actual": stk_c_vivo,
                "stock_destino_actual": stk_d_vivo,
                "stock_suficiente_cedis": stk_c_vivo >= float(r[8] or 0),
                "cliente": r[11] or '',
                "notas": r[12] or '',
                "solicitado_por": r[13],
                "vendedor": r[14] or '',
                "estatus": r[15],
                "autorizado_por": r[16] or '',
                "fecha_autorizacion": r[17] or '',
                "motivo_rechazo": r[18] or '',
                "folio_traspaso": r[19] or ''
            })

        return jsonify({"success": True, "solicitudes": solicitudes})
    except Exception as e:
        if conn_sq:
            try: conn_sq.close()
            except: pass
        return jsonify({"success": False, "solicitudes": [], "error": str(e)}), 500

@compras_bp.route('/api/compras/solicitudes-traspasos/<int:solicitud_id>/autorizar', methods=['POST'])
def autorizar_solicitud_traspaso(solicitud_id):
    if 'usuario' not in session:
        return jsonify({"success": False, "error": "No has iniciado sesión"}), 401

    usr = session.get('usuario', 'COMPRAS')
    ahora_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    conn = None
    try:
        conn = conectar_sqlite()
        cur = conn.cursor()
        cur.execute("SELECT id, folio, clave, cantidad, estatus FROM solicitudes_traspasos WHERE id = ?", (solicitud_id,))
        row = cur.fetchone()
        if not row:
            conn.close()
            return jsonify({"success": False, "error": "Solicitud no encontrada"}), 404

        if row[4] != 'PENDIENTE_COMPRAS':
            conn.close()
            return jsonify({"success": False, "error": f"La solicitud ya se encuentra con estatus {row[4]}"}), 400

        cur.execute("""
            UPDATE solicitudes_traspasos
            SET estatus = 'APROBADA',
                usuario_autoriza = ?,
                fecha_autorizacion = ?,
                motivo_rechazo = ''
            WHERE id = ?
        """, (usr, ahora_str, solicitud_id))
        conn.commit()
        conn.close()

        return jsonify({
            "success": True,
            "solicitud_id": solicitud_id,
            "folio": row[1],
            "mensaje": f"Solicitud {row[1]} autorizada exitosamente. Lista para surtido y preparación de traspaso en CEDIS."
        })
    except Exception as e:
        if conn:
            try: conn.close()
            except: pass
        return jsonify({"success": False, "error": str(e)}), 500

@compras_bp.route('/api/compras/solicitudes-traspasos/<int:solicitud_id>/rechazar', methods=['POST'])
def rechazar_solicitud_traspaso(solicitud_id):
    if 'usuario' not in session:
        return jsonify({"success": False, "error": "No has iniciado sesión"}), 401

    data = request.get_json() or {}
    motivo = str(data.get('motivo', '')).strip()
    if not motivo:
        return jsonify({"success": False, "error": "Debes especificar el motivo del rechazo"}), 400

    usr = session.get('usuario', 'COMPRAS')
    ahora_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    conn = None
    try:
        conn = conectar_sqlite()
        cur = conn.cursor()
        cur.execute("SELECT id, folio, clave, estatus FROM solicitudes_traspasos WHERE id = ?", (solicitud_id,))
        row = cur.fetchone()
        if not row:
            conn.close()
            return jsonify({"success": False, "error": "Solicitud no encontrada"}), 404

        if row[3] != 'PENDIENTE_COMPRAS':
            conn.close()
            return jsonify({"success": False, "error": f"La solicitud ya se encuentra con estatus {row[3]}"}), 400

        cur.execute("""
            UPDATE solicitudes_traspasos
            SET estatus = 'RECHAZADA',
                usuario_autoriza = ?,
                fecha_autorizacion = ?,
                motivo_rechazo = ?
            WHERE id = ?
        """, (usr, ahora_str, motivo, solicitud_id))
        conn.commit()
        conn.close()

        return jsonify({
            "success": True,
            "solicitud_id": solicitud_id,
            "folio": row[1],
            "mensaje": f"Solicitud {row[1]} rechazada correctamente. Motivo registrado: {motivo}"
        })
    except Exception as e:
        if conn:
            try: conn.close()
            except: pass
        return jsonify({"success": False, "error": str(e)}), 500

